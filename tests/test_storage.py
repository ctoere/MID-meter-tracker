"""Storage must never lose a row, reorder a column, or write in place."""

import openpyxl
import pytest

from app import schema, storage
from app.config import get_config


def test_round_trip_loses_nothing(sandbox):
    before = storage.load()
    counts = (len(before.chargers), len(before.meters), len(before.manifest),
              len(before.change_log), len(before.conflicts))

    storage.save(before)
    after = storage.load()

    assert (len(after.chargers), len(after.meters), len(after.manifest),
            len(after.change_log), len(after.conflicts)) == counts
    assert after.chargers == before.chargers
    assert after.meters == before.meters
    assert after.manifest == before.manifest


def test_round_trip_preserves_the_real_row_counts(sandbox):
    """The live register: 380 migrated + 67 imported chargers, 4 meters, 27 conflicts.

    A save must not change any of these. The numbers move only when an import
    or an edit deliberately moves them — update this test in the same commit.
    """
    before = storage.load()
    assert len(before.chargers) == 447
    assert len(before.meters) == 4
    assert len(before.conflicts) == 27
    storage.save(before)
    after = storage.load()
    assert (len(after.chargers), len(after.meters), len(after.conflicts)) == (447, 4, 27)


def test_column_order_is_stable(sandbox):
    def header(path, sheet):
        wb = openpyxl.load_workbook(path)
        row = [c.value for c in wb[sheet][1]] if sheet in wb.sheetnames else None
        wb.close()
        return row

    # A workbook predating a newly added sheet simply has no header for it yet;
    # save() creates it. Only sheets that already exist can be compared.
    before = {name: header(sandbox, name) for name in schema.MANAGED_SHEETS
              if header(sandbox, name) is not None}
    storage.save(storage.load())
    after = {name: header(sandbox, name) for name in before}

    assert after == before
    assert before[schema.SHEET_CHARGERS] == list(schema.CHARGER_COLUMNS)


def test_unmanaged_sheets_survive(sandbox):
    """'Read Me' is not ours to rewrite, and neither is anything a colleague adds."""
    wb = openpyxl.load_workbook(sandbox)
    wb["Read Me"]["A1"] = "Read Me content"
    extra = wb.create_sheet("Brand Coverage")
    extra.append(["Brand", "Notes"])
    extra.append(["Alfen", "hand-maintained"])
    wb.save(sandbox)
    wb.close()

    storage.save(storage.load())

    wb = openpyxl.load_workbook(sandbox)
    assert "Brand Coverage" in wb.sheetnames
    assert [c.value for c in wb["Brand Coverage"][2]] == ["Alfen", "hand-maintained"]
    assert wb["Read Me"]["A1"].value == "Read Me content"
    wb.close()


def test_extra_columns_are_not_dropped(sandbox):
    wb = openpyxl.load_workbook(sandbox)
    ws = wb[schema.SHEET_CHARGERS]
    ws.cell(row=1, column=ws.max_column + 1, value="Helpdesk Note")
    ws.cell(row=2, column=ws.max_column, value="ask Marijke")
    wb.save(sandbox)
    wb.close()

    storage.save(storage.load())
    after = storage.load()
    assert after.chargers[0]["Helpdesk Note"] == "ask Marijke"


def test_save_backs_up_first(sandbox):
    backup_dir = get_config().backup_dir
    assert not backup_dir.exists() or not list(backup_dir.glob("*.xlsx"))

    register = storage.load()
    register.chargers[0]["Notes (EN)"] = "changed"
    storage.save(register)

    backups = list(backup_dir.glob("*.xlsx"))
    assert len(backups) == 1
    # The backup holds the PREVIOUS content, not the new content.
    wb = openpyxl.load_workbook(backups[0])
    notes_column = list(schema.CHARGER_COLUMNS).index("Notes (EN)") + 1
    assert wb[schema.SHEET_CHARGERS].cell(row=2, column=notes_column).value != "changed"
    wb.close()


def test_backups_are_pruned_to_the_limit(sandbox):
    from app import config as config_module
    config_module.override_config(keep_backups=3)
    register = storage.load()
    for i in range(6):
        register.chargers[0]["Notes (EN)"] = f"revision {i}"
        storage.save(register)
    assert len(list(get_config().backup_dir.glob("*.xlsx"))) == 3


def test_write_is_atomic_and_leaves_no_temp_files(sandbox):
    storage.save(storage.load())
    leftovers = [p for p in sandbox.parent.iterdir() if p.suffix == ".tmp" or ".tmp." in p.name]
    assert leftovers == []


def test_failed_write_leaves_the_original_intact(sandbox, monkeypatch):
    original = sandbox.read_bytes()
    register = storage.load()

    def explode(self, path):
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", explode)
    with pytest.raises(OSError):
        storage.save(register)

    # Never written in place: the original survives a mid-write failure byte for byte.
    assert sandbox.read_bytes() == original
    assert [p for p in sandbox.parent.iterdir() if ".tmp." in p.name] == []


def test_refuses_to_clobber_an_edit_made_behind_our_back(sandbox):
    register = storage.load()
    register.loaded_mtime = 0.0  # simulate someone saving from Excel after we loaded
    with pytest.raises(storage.StaleRegisterError):
        storage.save(register)
    storage.save(register, force=True)  # explicit override still works


def test_a_sheet_added_after_the_workbook_was_created_is_written_on_save(sandbox):
    """Upgrading an older register must not need a migration step."""
    wb = openpyxl.load_workbook(sandbox)
    if schema.SHEET_CONFORMITY in wb.sheetnames:
        del wb[schema.SHEET_CONFORMITY]
        wb.save(sandbox)
    wb.close()

    register = storage.load()          # missing sheet reads as empty, not an error
    assert register.conformity == []
    storage.save(register)

    wb = openpyxl.load_workbook(sandbox)
    assert schema.SHEET_CONFORMITY in wb.sheetnames
    assert [c.value for c in wb[schema.SHEET_CONFORMITY][1]] == list(schema.CONFORMITY_COLUMNS)
    wb.close()
